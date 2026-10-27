"""
The sales report as an .xlsx workbook, for a period and one showroom or all of them:

  Summary   per showroom: bills, units, sales, GST, returns and net sales (with cost, profit and margin for the admin)
  Bills     one row per bill, with how it was paid
  Items     one row per item sold
  Returns   the credit notes (returns and exchanges)
  Cash      the daily cash closes and their differences

Cost and profit columns are only in the admin's report.
"""

from collections import defaultdict
from collections.abc import Sequence
from datetime import date, datetime, time, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core import clock
from app.modules.auth.access import visible_location
from app.modules.cash.models import CashClose
from app.modules.locations.constants import SHOWROOMS
from app.modules.returns.models import SaleReturn
from app.modules.sales.models import Sale
from app.modules.users.models import User

HEAD = Font(bold=True, color='FFFFFF')
HEAD_FILL = PatternFill('solid', fgColor='175C26')
MONEY = '#,##0.00'
CITY = {s.id: s.city for s in SHOWROOMS}
METHOD = {'cash': 'Cash', 'upi': 'UPI', 'card': 'Card', 'split': 'Split', 'credit': 'Exchange credit'}


def _bounds(since: date | None, until: date | None) -> tuple[datetime | None, datetime | None]:
    start = datetime.combine(since, time(0), clock.IST) if since else None
    end = datetime.combine(until, time(0), clock.IST) + timedelta(days=1) if until else None
    return start, end


def _local(at: datetime) -> datetime:
    return at.astimezone(clock.IST).replace(tzinfo=None)  # Excel has no time zones


def _sheet(wb: Workbook, title: str, header: Sequence[str], rows: list[list], money: Sequence[str] = ()) -> None:
    ws = wb.create_sheet(title)
    ws.append(list(header))
    for cell in ws[1]:
        cell.font = HEAD
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(vertical='center')
    for row in rows:
        ws.append(row)
    for i, name in enumerate(header, start=1):
        letter = get_column_letter(i)
        width = max([len(str(name))] + [len(str(r[i - 1])) for r in rows[:500] if r[i - 1] is not None])
        ws.column_dimensions[letter].width = min(max(width + 2, 8), 48)
        if name in money:
            for cell in ws[letter][1:]:
                cell.number_format = MONEY
        if rows and isinstance(rows[0][i - 1], datetime):
            for cell in ws[letter][1:]:
                cell.number_format = 'dd-mmm-yyyy hh:mm'
        elif rows and isinstance(rows[0][i - 1], date):
            for cell in ws[letter][1:]:
                cell.number_format = 'dd-mmm-yyyy'
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions


def _paid(p: dict) -> tuple[float, float, float]:
    """Cash, UPI and card parts of a payment."""
    method = p.get('method')
    if method == 'cash':
        return p.get('received', 0) - p.get('change', 0), 0, 0
    if method == 'upi':
        return 0, p.get('received', 0), 0
    if method == 'card':
        return 0, 0, p.get('received', 0)
    if method == 'split' and p.get('split'):
        return p['split'].get('cash', 0), p['split'].get('upi', 0), 0
    return 0, 0, 0


def build(
    session: Session, user: User, *, location: str | None, since: date | None, until: date | None
) -> tuple[bytes, str]:
    """The workbook's bytes and a file name such as sales-chennai-2026-11-01-to-2026-11-30.xlsx."""
    admin = user.is_admin
    location_id = visible_location(user, location)
    start, end = _bounds(since, until)

    stmt = select(Sale).order_by(Sale.at)
    rstmt = select(SaleReturn).order_by(SaleReturn.at)
    cstmt = select(CashClose).order_by(CashClose.day, CashClose.location_id)
    if location_id:
        stmt = stmt.where(Sale.location_id == location_id)
        rstmt = rstmt.where(SaleReturn.location_id == location_id)
        cstmt = cstmt.where(CashClose.location_id == location_id)
    if start:
        stmt = stmt.where(Sale.at >= start)
        rstmt = rstmt.where(SaleReturn.at >= start)
        cstmt = cstmt.where(CashClose.day >= since)
    if end:
        stmt = stmt.where(Sale.at < end)
        rstmt = rstmt.where(SaleReturn.at < end)
        cstmt = cstmt.where(CashClose.day <= until)
    sales = session.scalars(stmt).all()
    returns = session.scalars(rstmt).unique().all()
    closes = session.scalars(cstmt).all()

    wb = Workbook()
    wb.remove(wb.active)

    # Summary
    per = defaultdict(lambda: defaultdict(float))
    for s in sales:
        t = s.totals
        row = per[s.location_id]
        row['bills'] += 1
        row['units'] += t['qty']
        row['sales'] += t['payable']
        row['taxable'] += t['taxable']
        row['gst'] += t['tax']
        row['cost'] += sum(float(ln.cost) * ln.qty for ln in s.lines)
    for r in returns:
        per[r.location_id]['returns'] += float(r.amount)
        per[r.location_id]['returned_cost'] += sum(
            float(next((ln.cost for ln in r.sale.lines if ln.code == x['code']), 0)) * x['qty'] for x in r.lines
        )
    header = ['Showroom', 'Bills', 'Units', 'Sales (incl. GST)', 'GST', 'Returns', 'Net sales']
    money = ['Sales (incl. GST)', 'GST', 'Returns', 'Net sales']
    if admin:
        header += ['Purchase cost', 'Gross profit', 'Margin %']
        money += ['Purchase cost', 'Gross profit']
    rows = []
    shops = [s.id for s in SHOWROOMS if not location_id or s.id == location_id]
    for shop in shops:
        p = per.get(shop, defaultdict(float))
        row = [
            CITY[shop],
            int(p['bills']),
            int(p['units']),
            p['sales'],
            p['gst'],
            p['returns'],
            p['sales'] - p['returns'],
        ]
        if admin:
            # profit on what stayed sold: the returned share's GST-free value and cost come off
            ratio = (p['sales'] - p['returns']) / p['sales'] if p['sales'] else 0
            cost = p['cost'] - p['returned_cost']
            profit = p['taxable'] * ratio - cost
            row += [cost, profit, round(profit / (p['taxable'] * ratio) * 100, 1) if p['taxable'] * ratio else 0]
        rows.append(row)
    if len(shops) > 1:
        total = ['All showrooms'] + [sum(r[i] for r in rows) for i in range(1, len(header))]
        if admin:
            base = sum(
                per[s]['taxable'] * ((per[s]['sales'] - per[s]['returns']) / per[s]['sales'] if per[s]['sales'] else 0)
                for s in shops
            )
            total[-1] = round(total[-2] / base * 100, 1) if base else 0
        rows.append(total)
    _sheet(wb, 'Summary', header, rows, money)

    # Bills
    header = [
        'Bill no',
        'Date',
        'Showroom',
        'Customer',
        'Mobile',
        'Items',
        'Qty',
        'MRP total',
        'Savings',
        'Taxable',
        'GST',
        'Round-off',
        'Amount',
        'Paid by',
        'Cash',
        'UPI',
        'Card',
        'Exchange credit',
    ]
    money = ['MRP total', 'Savings', 'Taxable', 'GST', 'Round-off', 'Amount', 'Cash', 'UPI', 'Card', 'Exchange credit']
    if admin:
        header += ['Purchase cost', 'Gross profit']
        money += ['Purchase cost', 'Gross profit']
    rows = []
    for s in sales:
        t, p = s.totals, s.payment
        cash, upi, card = _paid(p)
        row = [
            s.bill_no,
            _local(s.at),
            CITY.get(s.location_id, s.location_id),
            s.name,
            s.phone,
            t['count'],
            t['qty'],
            t['mrpTotal'],
            t['savings'],
            t['taxable'],
            t['tax'],
            t['roundOff'],
            t['payable'],
            METHOD.get(p.get('method'), p.get('method')),
            cash,
            upi,
            card,
            p.get('credit', 0),
        ]
        if admin:
            cost = sum(float(ln.cost) * ln.qty for ln in s.lines)
            row += [cost, t['taxable'] - cost]
        rows.append(row)
    _sheet(wb, 'Bills', header, rows, money)

    # Items
    header = ['Bill no', 'Date', 'Showroom', 'Code', 'Product', 'Qty', 'MRP', 'Price', 'Amount', 'GST %']
    money = ['MRP', 'Price', 'Amount']
    if admin:
        header += ['Cost', 'Cost total']
        money += ['Cost', 'Cost total']
    rows = []
    for s in sales:
        for ln in s.lines:
            row = [
                s.bill_no,
                _local(s.at),
                CITY.get(s.location_id, s.location_id),
                ln.code,
                ln.name,
                ln.qty,
                float(ln.mrp),
                float(ln.price),
                float(ln.price) * ln.qty,
                float(ln.gst),
            ]
            if admin:
                row += [float(ln.cost), float(ln.cost) * ln.qty]
            rows.append(row)
    _sheet(wb, 'Items', header, rows, money)

    # Returns
    header = [
        'Credit note',
        'Date',
        'Showroom',
        'Against bill',
        'Kind',
        'Customer',
        'Items',
        'Value',
        'Cash refund',
        'New bill',
        'Reason',
    ]
    rows = [
        [
            r.return_no,
            _local(r.at),
            CITY.get(r.location_id, r.location_id),
            r.sale.bill_no,
            r.kind.title(),
            r.sale.name,
            ', '.join(f'{x["code"]} × {x["qty"]}' for x in r.lines),
            float(r.amount),
            float(r.refund),
            r.exchange_sale.bill_no if r.exchange_sale else '',
            r.reason,
        ]
        for r in returns
    ]
    _sheet(wb, 'Returns', header, rows, ['Value', 'Cash refund'])

    # Cash
    header = [
        'Day',
        'Showroom',
        'Opening float',
        'Cash from bills',
        'Cash refunds',
        'Expected',
        'Counted',
        'Difference',
        'Kept for next day',
        'Closed by',
        'Note',
    ]
    rows = [
        [
            c.day,
            CITY.get(c.location_id, c.location_id),
            float(c.opening),
            float(c.cash_sales),
            float(c.cash_refunds),
            float(c.expected),
            float(c.counted),
            float(c.difference),
            float(c.float_next),
            c.by,
            c.note,
        ]
        for c in closes
    ]
    _sheet(
        wb,
        'Cash',
        header,
        rows,
        ['Opening float', 'Cash from bills', 'Cash refunds', 'Expected', 'Counted', 'Difference', 'Kept for next day'],
    )

    out = BytesIO()
    wb.save(out)
    span = f'{since or "start"}-to-{until or clock.now().astimezone(clock.IST).date()}'
    return out.getvalue(), f'sales-{location_id or "all-showrooms"}-{span}.xlsx'
