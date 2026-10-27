"""
Builds the demo schema from nothing: the tables, the showrooms and logins, the catalogue from the product Excel,
and a month of sample activity made through the same services the app uses, so every figure adds up:

  opening stock   each showroom's sheet in seed/opening-stock-2026, imported 5 weeks ago
  a delivery      a received purchase order of the season's products per showroom, a month ago
  sales           about 30 days of bills, busier towards Christmas, only for what is in stock
  returns         a few returns and one exchange
  cash closes     every past day closed, a few with a small difference
  orders          one of each state: to count in, to review, rejected, cancelled
  quotations      a couple per showroom, one of them expired
  stock returns   one on its way to the admin, one counted in

It is made the same way every time (a seeded random generator), relative to now.
"""

import math
import random
import uuid
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.orm import Session

from app import db_models  # noqa: F401  (registers every table)
from app.core import clock
from app.core.config import BACKEND_DIR
from app.core.database import DEMO_SCHEMA, Base, get_demo_engine, get_engine, get_session_factory
from app.modules.cash import service as cash
from app.modules.cash.schemas import CloseIn
from app.modules.locations import repository as locations
from app.modules.locations.constants import ADMIN_USERNAME, SHOWROOMS
from app.modules.orders import service as orders
from app.modules.orders.schemas import CountedLineIn, OrderEditIn, OrderIn, OrderLineIn, ReceiveIn
from app.modules.products import repository as products
from app.modules.products import service as products_service
from app.modules.products.importer import DEFAULT_EXCEL, read_catalogue
from app.modules.quotes import service as quotes
from app.modules.quotes.schemas import QuoteIn
from app.modules.returns import service as returns
from app.modules.returns.schemas import ExchangeIn, ReturnIn
from app.modules.sales import bill
from app.modules.sales import service as sales
from app.modules.sales.schemas import LineIn, PaymentIn, SaleIn, SplitIn
from app.modules.stock import repository as stock
from app.modules.stock import service as stock_service
from app.modules.stock.schemas import ImportLineIn, StockImportIn
from app.modules.transfers import service as transfers
from app.modules.transfers.schemas import StockReturnIn, StockReturnReceiveIn
from app.modules.users import repository as users
from app.modules.users import service as users_service

OPENING_DIR = BACKEND_DIR / 'seed' / 'opening-stock-2026'
DAYS = 30
FLOAT = 2000  # cash kept in each drawer overnight
WEIGHT = {'chennai': 1.4, 'bengaluru': 1.1, 'mumbai': 1.2, 'pune': 0.8, 'goa': 0.6}
METHODS = ['cash', 'cash', 'upi', 'upi', 'upi', 'card', 'split']
NAMES = [
    'Priya Raman',
    'Arjun Mehta',
    'Kavya Iyer',
    'Rohan Desai',
    'Ananya Rao',
    'Vikram Nair',
    'Sneha Kulkarni',
    'Aditya Shetty',
    'Meera Pillai',
    'Karthik Subramanian',
    'Divya Menon',
    'Rahul Joshi',
    'Lakshmi Narayan',
    'Siddharth Fernandes',
    'Pooja Gupta',
    'Nikhil Patil',
    'Shreya Das',
    'Joseph DSouza',
    'Anjali Krishnan',
    'Farhan Sheikh',
    'Neha Bhat',
    'Varun Reddy',
    'Maria Pereira',
    'Ishaan Kapoor',
    'Ritika Sharma',
]


@dataclass(frozen=True, slots=True)
class DemoReport:
    sales: int
    returns: int
    orders: int
    quotes: int

    def __str__(self) -> str:
        return (
            f'Demo ready: 5 showrooms and the admin, {self.sales} sales, {self.returns} returns, '
            f'{self.orders} purchase orders and {self.quotes} quotations.'
        )


def demo_exists() -> bool:
    with get_engine().connect() as conn:
        return bool(
            conn.scalar(
                text("select 1 from information_schema.tables where table_schema = :s and table_name = 'users'"),
                {'s': DEMO_SCHEMA},
            )
        )


def _recreate_schema() -> None:
    with get_engine().begin() as conn:
        conn.execute(text(f'DROP SCHEMA IF EXISTS {DEMO_SCHEMA} CASCADE'))
        conn.execute(text(f'CREATE SCHEMA {DEMO_SCHEMA}'))
    with get_demo_engine().begin() as conn:
        Base.metadata.create_all(conn)


def _opening(path: Path) -> list[tuple[str, int, float]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(wb.active.iter_rows(values_only=True))
    finally:
        wb.close()
    header = [str(h or '').strip().lower() for h in rows[0]]
    code, qty, rate = header.index('code'), header.index('qty'), header.index('purchase rate')
    return [(str(r[code]).strip().upper(), int(r[qty]), float(r[rate] or 0)) for r in rows[1:] if r[code] and r[qty]]


def _at(day: date, rand: random.Random, start: int = 10, end: int = 20) -> datetime:
    return datetime.combine(day, time(rand.randint(start, end - 1), rand.randint(0, 59)), clock.IST)


def _mobile(rand: random.Random) -> str:
    return f'{rand.choice("6789")}{rand.randint(0, 999_999_999):09d}'


class _Builder:
    def __init__(self, session: Session, now: datetime) -> None:
        self.s = session
        self.now = now
        self.today = now.astimezone(clock.IST).date()
        self.rand = random.Random(2026)
        self.admin = users.get(session, ADMIN_USERNAME)
        self.shops = {sh.id: users.get(session, sh.id) for sh in SHOWROOMS}
        self.counts = {'sales': 0, 'returns': 0, 'orders': 0, 'quotes': 0}
        self.bills: dict[str, list] = {sh.id: [] for sh in SHOWROOMS}

    def day(self, ago: int) -> date:
        return self.today - timedelta(days=ago)

    def have(self, location: str) -> dict[str, int]:
        return {lv.code: lv.qty for lv in stock.levels(self.s, location) if lv.qty > 0}

    def opening_stock(self) -> None:
        for sh in SHOWROOMS:
            path = OPENING_DIR / f'opening-stock-{sh.id}.xlsx'
            rows = _opening(path) if path.exists() else []
            known = products.by_codes(self.s, (c for c, _, _ in rows))
            lines = [ImportLineIn(code=c, qty=q, rate=r) for c, q, r in rows if c in known]
            if lines:
                stock_service.import_stock(
                    self.s,
                    self.shops[sh.id],
                    StockImportIn(id=uuid.uuid4(), file=path.name, lines=lines),
                    at=_at(self.day(DAYS + 5), self.rand),
                )

    def season_delivery(self) -> None:
        season = [p for p in products.list_all(self.s) if p.in_season]
        for sh in SHOWROOMS:
            user = self.shops[sh.id]
            lines = [OrderLineIn(code=p.code, qty=self.rand.randint(12, 30)) for p in season]
            order = orders.create_order(
                self.s, user, OrderIn(lines=lines, note='Season stock'), at=_at(self.day(DAYS + 3), self.rand)
            )
            self.counts['orders'] += 1
            orders.send(self.s, self.admin, order.id, OrderEditIn(lines=lines), at=_at(self.day(DAYS + 2), self.rand))
            counted = [CountedLineIn(code=ln.code, qty=ln.qty) for ln in lines]
            counted[0] = CountedLineIn(code=lines[0].code, qty=max(lines[0].qty - 2, 1))  # a little short
            orders.receive(
                self.s,
                user,
                order.id,
                ReceiveIn(counted=counted, note='2 short in the first carton'),
                at=_at(self.day(DAYS + 1), self.rand),
            )

    def payment(self, payable: float) -> PaymentIn:
        method = self.rand.choice(METHODS)
        if method == 'cash':
            step = 500 if payable > 2000 else 100
            return PaymentIn(method='cash', received=math.ceil(payable / step) * step)
        if method == 'split':
            part = max(100, math.floor(payable * self.rand.uniform(0.2, 0.6) / 100) * 100)
            if part >= payable:
                return PaymentIn(method='upi')
            return PaymentIn(method='split', split=SplitIn(cash=part, upi=payable - part))
        if method == 'card':
            return PaymentIn(method='card', approval=f'{self.rand.randint(100000, 999999)}')
        return PaymentIn(method='upi')

    def sell(self) -> None:
        for ago in range(DAYS - 1, -1, -1):
            rush = 1 + (DAYS - ago) / DAYS
            day = self.day(ago)
            for sh in SHOWROOMS:
                have = self.have(sh.id)
                count = round((2 + self.rand.random() * 4) * WEIGHT[sh.id] * rush)
                # through the day in order, so the bill numbers run with the time
                for at in sorted(_at(day, self.rand) for _ in range(count)):
                    if at > self.now or not have:
                        continue
                    items: dict[str, int] = {}
                    for _ in range(self.rand.randint(1, 4)):
                        code = self.rand.choice(sorted(have))
                        if have[code] - items.get(code, 0) > 0:
                            items[code] = items.get(code, 0) + 1
                    if not items:
                        continue
                    for code, q in items.items():
                        have[code] -= q
                        if have[code] <= 0:
                            del have[code]
                    catalogue = products.by_codes(self.s, items)
                    payable = bill.totals(bill.bill_lines(items.items(), catalogue))['payable']
                    sale = sales.create_sale(
                        self.s,
                        self.shops[sh.id],
                        SaleIn(
                            id=uuid.uuid4(),
                            lines=[LineIn(code=c, qty=q) for c, q in items.items()],
                            phone=_mobile(self.rand),
                            name=self.rand.choice(NAMES),
                            payment=self.payment(payable),
                        ),
                        at=at,
                    )
                    self.bills[sh.id].append(sale)
                    self.counts['sales'] += 1

    def take_returns(self) -> None:
        for i, sh in enumerate(SHOWROOMS):
            recent = [b for b in self.bills[sh.id] if (self.now - b.at).days <= 8]
            if not recent:
                continue
            sale = recent[len(recent) // 2]
            line = sale.lines[0]
            at = min(sale.at + timedelta(days=1, hours=2), self.now - timedelta(minutes=30))
            exchange = None
            if i == 0:  # Chennai: one exchange, for another product in stock
                have = self.have(sh.id)
                other = next((c for c in sorted(have) if c != line.code), None)
                if other:
                    exchange = ExchangeIn(
                        id=uuid.uuid4(), lines=[LineIn(code=other, qty=1)], payment=PaymentIn(method='upi')
                    )
            returns.create_return(
                self.s,
                self.shops[sh.id],
                ReturnIn(
                    id=uuid.uuid4(),
                    sale_id=sale.id,
                    lines=[LineIn(code=line.code, qty=1)],
                    reason='Wanted a different size' if exchange else 'Gift not needed',
                    exchange=exchange,
                ),
                at=at,
            )
            self.counts['returns'] += 1

    def close_days(self) -> None:
        for sh in SHOWROOMS:
            opening = Decimal(FLOAT)
            for ago in range(DAYS - 1, 0, -1):
                day = self.day(ago)
                f = cash.get_day(self.s, self.shops[sh.id], None, day)
                expected = Decimal(str(opening)) + Decimal(str(f.cash_sales)) - Decimal(str(f.cash_refunds))
                off = self.rand.choice([0] * 8 + [-200, -50, 100])
                counted = max(expected + off, Decimal(0))
                cash.close_day(
                    self.s,
                    self.shops[sh.id],
                    CloseIn(
                        day=day,
                        opening=float(opening),
                        counted=float(counted),
                        float_next=float(min(FLOAT, counted)),
                        note='' if not off else ('Short: checking with the counter staff' if off < 0 else 'Over'),
                    ),
                    at=datetime.combine(day, time(21, 15), clock.IST),
                )
                opening = min(Decimal(FLOAT), counted)

    def open_orders(self) -> None:
        def lines(location: str, n: int) -> list[OrderLineIn]:
            low = sorted(self.have(location).items(), key=lambda kv: kv[1])[:n]
            return [OrderLineIn(code=c, qty=self.rand.randint(6, 15)) for c, _ in low] or [
                OrderLineIn(code=p.code, qty=10) for p in products.list_all(self.s)[:n]
            ]

        def request(location: str, ago: int, note: str):
            self.counts['orders'] += 1
            return orders.create_order(
                self.s,
                self.shops[location],
                OrderIn(lines=lines(location, 4), note=note),
                at=_at(self.day(ago), self.rand),
            )

        o = request('chennai', 3, 'Running low before the weekend')
        orders.send(
            self.s,
            self.admin,
            o.id,
            OrderEditIn(lines=o_lines(o), note='Sent with the Monday van'),
            at=_at(self.day(2), self.rand),
        )
        request('bengaluru', 1, 'Trees are selling fast')
        request('mumbai', 0, '')
        o = request('goa', 4, '')
        orders.reject(self.s, self.admin, o.id, 'Goa has enough of these; ask Mumbai', at=_at(self.day(3), self.rand))
        o = request('pune', 2, '')
        orders.cancel(self.s, self.shops['pune'], o.id, at=_at(self.day(2), self.rand))

    def write_quotes(self) -> None:
        season = [p.code for p in products.list_all(self.s) if p.in_season]
        for sh in SHOWROOMS:
            for ago, days in ((12, 7), (2, 15)):
                codes = self.rand.sample(season, k=min(3, len(season)))
                quotes.create_quote(
                    self.s,
                    self.shops[sh.id],
                    QuoteIn(
                        id=uuid.uuid4(),
                        lines=[LineIn(code=c, qty=self.rand.randint(1, 3)) for c in codes],
                        name=self.rand.choice(NAMES),
                        phone=_mobile(self.rand),
                        valid_days=days,
                        note='Corporate order: delivery before 20 December' if ago == 2 else '',
                    ),
                    at=_at(self.day(ago), self.rand),
                )
                self.counts['quotes'] += 1

    def send_back(self) -> None:
        for location, ago, done in (('goa', 6, True), ('pune', 1, False)):
            have = sorted(self.have(location).items(), key=lambda kv: -kv[1])[:2]
            if not have:
                continue
            r = transfers.send_back(
                self.s,
                self.shops[location],
                StockReturnIn(
                    id=uuid.uuid4(),
                    lines=[LineIn(code=c, qty=min(q, 3)) for c, q in have],
                    reason='another-store' if done else 'end-of-season',
                    note='For Mumbai' if done else '',
                ),
                at=_at(self.day(ago), self.rand),
            )
            if done:
                counted = [CountedLineIn(code=ln['code'], qty=ln['qty']) for ln in r.lines]
                transfers.receive(
                    self.s,
                    self.admin,
                    r.id,
                    StockReturnReceiveIn(counted=counted),
                    at=_at(self.day(ago - 2), self.rand),
                )


def o_lines(order) -> list[OrderLineIn]:
    return [OrderLineIn(code=ln['code'], qty=ln['qty'], rate=ln['rate']) for ln in order.lines]


def build_demo(excel: Path = DEFAULT_EXCEL, *, now: datetime | None = None) -> DemoReport:
    """Throws away the demo schema, if there is one, and builds it again."""
    catalogue = read_catalogue(excel)
    _recreate_schema()
    with get_session_factory(demo=True)() as session:
        locations.add_missing_showrooms(session)
        users_service.add_missing_logins(session)
        products_service.import_catalogue(session, catalogue, season_from_excel=True)
        session.commit()

        b = _Builder(session, now or clock.now())
        b.opening_stock()
        b.season_delivery()
        b.sell()
        b.take_returns()
        b.close_days()
        b.open_orders()
        b.write_quotes()
        b.send_back()
        return DemoReport(**b.counts)
