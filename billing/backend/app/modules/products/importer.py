"""
Reads the product Excel (seed/products-clean.xlsx, made from the old billing software's export by
seed/clean_products.py) into plain rows for the products table. No database here.
"""

from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from typing import TypedDict

from openpyxl import load_workbook

from app.core.config import BACKEND_DIR

DEFAULT_EXCEL = BACKEND_DIR / 'seed' / 'products-clean.xlsx'
SHEET = 'Products'
SEASON_PREFIX = 'In Season'  # the season column is "In Season 2026": its suffix names the season
REQUIRED_COLUMNS = (
    'Code',
    'Name',
    'Type',
    'Department',
    'Barcode',
    'HSN',
    'MRP',
    'Offer Price',
    'Cost Price',
    'GST % (Sale)',
)
YES = frozenset({'Y', 'YES', '1', 'TRUE'})


class ProductRow(TypedDict):
    code: str
    name: str
    type: str
    dept: str
    barcode: str
    hsn: str
    mrp: Decimal
    offer: Decimal
    cost: Decimal
    gst: Decimal
    in_season: bool


@dataclass(frozen=True, slots=True)
class ImportedCatalogue:
    season: str
    products: list[ProductRow]


class CatalogueFileError(ValueError):
    """The Excel isn't in the expected shape: no Products sheet, a missing column, or a code used twice."""


def _text(value: object) -> str:
    return '' if value is None else str(value).strip()


def _money(value: object) -> Decimal:
    return Decimal(str(value)) if isinstance(value, int | float) and not isinstance(value, bool) else Decimal(0)


def sentence_case(text: str) -> str:
    """Departments as the app shows them: TREE HANGINGS -> Tree hangings."""
    return text[:1].upper() + text[1:].lower()


def read_catalogue(path: Path = DEFAULT_EXCEL) -> ImportedCatalogue:
    """The products in the Excel, and the season named in its "In Season ..." column header."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if SHEET not in workbook.sheetnames:
            raise CatalogueFileError(f'{path.name} has no "{SHEET}" sheet.')
        rows = workbook[SHEET].iter_rows(values_only=True)
        header = [_text(h) for h in next(rows, ())]
        missing = [c for c in REQUIRED_COLUMNS if c not in header]
        season_col = next((h for h in header if h.startswith(SEASON_PREFIX)), None)
        if missing or season_col is None:
            raise CatalogueFileError(f'{path.name} is missing columns: {", ".join(missing or [SEASON_PREFIX])}.')

        products: list[ProductRow] = []
        seen: set[str] = set()
        for values in rows:
            r = dict(zip(header, values, strict=False))
            code = _text(r['Code']).upper()
            if not code:
                continue
            if code in seen:
                raise CatalogueFileError(f'{path.name} has the code {code} more than once.')
            seen.add(code)
            products.append(
                ProductRow(
                    code=code,
                    name=_text(r['Name']),
                    type=_text(r['Type']) or 'Product',
                    dept=sentence_case(_text(r['Department'])),
                    barcode=_text(r['Barcode']),
                    hsn=_text(r['HSN']),
                    mrp=_money(r['MRP']),
                    offer=_money(r['Offer Price']),
                    cost=_money(r['Cost Price']),
                    gst=_money(r['GST % (Sale)']),
                    in_season=_text(r[season_col]).upper() in YES,
                )
            )
    finally:
        workbook.close()
    return ImportedCatalogue(season=season_col.removeprefix(SEASON_PREFIX).strip(), products=products)
