"""
Clean a product export from the old billing software into a new workbook.

    python clean_products.py products.xlsx products-clean.xlsx

Needs: pip install xlrd openpyxl. (products.xlsx is really an old .xls file, which only xlrd reads.)
Run it only on a fresh export: it rewrites products-clean.xlsx, including the "In Season 2026" column.
"""
import sys, re
import xlrd
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

SRC, OUT = sys.argv[1], sys.argv[2]
book = xlrd.open_workbook(SRC)
sh = book.sheet_by_index(0)
H = [h.strip() for h in sh.row_values(2)]
col = {h: i for i, h in enumerate(H)}


def txt(v):
    """Cell value as clean text: floats like 850043.0 -> '850043', whitespace collapsed."""
    if v == '' or v is None:
        return ''
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else str(v)
    return re.sub(r'[ \t]+', ' ', str(v)).strip()


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


MONTHS = 'Jan Feb Mar Apr May Jun Jul Aug Sep Oct Nov Dec'.split()


def added(r):
    """Mfr.Date -> (year, month, issue). Typed 'm/yy' was read by Excel as day=yy in 2026."""
    c = col['Mfr.Date']
    t, v = sh.cell_type(r, c), sh.cell_value(r, c)
    if t == xlrd.XL_CELL_EMPTY or v == '':
        return None, None, None
    if t == xlrd.XL_CELL_DATE:
        d = xlrd.xldate_as_datetime(v, book.datemode)
        return 2000 + d.day, d.month, None  # 2026-09-19 -> Sep 2019
    if t == xlrd.XL_CELL_NUMBER:
        y = int(v)
        if 2010 <= y <= 2026:
            return y, None, None
        return None, None, f'Unclear added date "{y}"'
    s = str(v).strip()
    m = re.fullmatch(r'(\d{1,2})[\\/](20\d\d)', s)
    if m:
        return int(m.group(2)), int(m.group(1)), None
    m = re.fullmatch(r'/(\d\d)', s)
    if m:
        return 2000 + int(m.group(1)), None, None
    return None, None, f'Unclear added date "{s}"'


products, review, removed = [], [], []
for r in range(3, sh.nrows):
    raw = sh.row_values(r)

    def g(h):
        return raw[col[h]]

    code, name = txt(g('Code')), txt(g('Product Name'))
    source_row = r + 1  # Excel row number in the original file

    if name == '' or name.lower() == 'fsa':
        why = 'Blank name (junk row)' if name == '' else 'Looks like a test entry ("fsa")'
        removed.append([source_row, code, name, txt(g('Department')), why])
        continue

    issues = []
    if '\n' in name or name.startswith('"'):
        issues.append('Name is a description, needs a proper product name')
        name = re.sub(r'\s+', ' ', name.strip('"')).strip()

    mrp, cost, offer = num(g('Ret.Rate')), num(g('D.Rate')), num(g('S.Price'))
    if mrp == 0:
        issues.append('MRP is 0')
    if cost == 0:
        issues.append('Cost (D.Rate) is 0')
    if offer and mrp and offer > mrp:
        issues.append('Offer price is higher than MRP')
    disc = round((1 - offer / mrp) * 100, 2) if offer and mrp else None
    if disc is not None and abs(disc - num(g('DiscP'))) >= 0.1:
        issues.append(f'Discount % corrected ({num(g("DiscP")):g} -> {disc:g})')

    bc_raw = g('Supl Code')
    barcode = txt(bc_raw)
    if isinstance(bc_raw, float) and len(barcode) >= 12 and re.search(r'0{5,}$', barcode):
        issues.append(f'Barcode damaged by Excel ({barcode}), re-scan needed')
        barcode = ''

    y, m, date_issue = added(r)
    if date_issue:
        issues.append(date_issue)

    gst = num(g('S.Tax%'))
    upper = name.upper()
    kind = 'Service' if ('COURIER' in upper or 'TRANSPORT' in upper) else 'Product'
    if upper == 'CORIANDER':
        issues.append('Not a Christmas product? (0% GST)')
    elif gst == 0 and kind == 'Product':
        issues.append('GST is 0%')

    row = {
        'Code': code, 'Name': name, 'Type': kind,
        'Department': txt(g('Department')), 'Dept No': int(num(g('Dept'))),
        'Sub Code': txt(g('SCode')).upper(),
        'Supplier Item No': txt(g('Product No.')), 'Barcode': barcode, 'HSN': txt(g('HSN')),
        'MRP': mrp, 'Offer Price': offer or None, 'Discount %': disc, 'Cost Price': cost,
        'GST % (Sale)': gst, 'Tax % (Purchase)': num(g('P.Tax%')),
        'Manufacturer': txt(g('Mfr.Name')), 'Origin': txt(g('Mfr.Add.')),
        'Added Year': y, 'Added Month': MONTHS[m - 1] if m else None,
        'In Season 2026': '',
        'Source Row': source_row,
    }
    products.append(row)
    if issues:
        review.append([source_row, code, name, row['Department'], '; '.join(issues)])

# ---------- write ----------
wb = Workbook()
HEAD = Font(bold=True, color='FFFFFF')
FILL = PatternFill('solid', fgColor='1F4E3D')
FILL_INPUT_HEAD = PatternFill('solid', fgColor='BF8F00')
FILL_INPUT = PatternFill('solid', fgColor='FFF2CC')


def sheet(ws, headers, rows, widths, money=(), pct=(), text=(), inputs=()):
    ws.append(headers)
    for c in ws[1]:
        c.font = HEAD
        c.fill = FILL_INPUT_HEAD if c.value in inputs else FILL
    for rw in rows:
        ws.append(rw)
    ws.freeze_panes = 'C2' if len(headers) > 6 else 'A2'
    ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(headers, 1):
        L = get_column_letter(i)
        ws.column_dimensions[L].width = widths.get(h, 14)
        fmt = '#,##0.00' if h in money else '0.00' if h in pct else '@' if h in text else None
        for c in ws[L][1:]:
            if fmt:
                c.number_format = fmt
            if h in inputs:
                c.fill = FILL_INPUT


P_HEAD = list(products[0].keys())
ws = wb.active
ws.title = 'Products'
sheet(ws, P_HEAD, [[p[h] for h in P_HEAD] for p in products],
      {'Code': 10, 'Name': 42, 'Department': 18, 'Supplier Item No': 18, 'Barcode': 16, 'Manufacturer': 30,
       'In Season 2026': 15, 'Type': 9, 'Dept No': 8, 'Sub Code': 11, 'HSN': 10},
      money=('MRP', 'Offer Price', 'Cost Price'), pct=('Discount %',),
      text=('Code', 'Supplier Item No', 'Barcode', 'HSN'), inputs=('In Season 2026',))

R_HEAD = ['Source Row', 'Code', 'Name', 'Department', 'What to check']
sheet(wb.create_sheet('Needs Review'), R_HEAD, review, {'Name': 42, 'Department': 18, 'What to check': 80})
sheet(wb.create_sheet('Removed'), R_HEAD[:4] + ['Why removed'], removed,
      {'Name': 30, 'Department': 18, 'Why removed': 40})

notes = wb.create_sheet('Notes')
NOTES = [
    ('Column', 'Original column', 'Meaning / what changed'),
    ('Code', 'Code', 'Product ID (unique). Stored as text.'),
    ('Name', 'Product Name', 'Extra spaces removed. Original casing kept.'),
    ('Type', '(new)', 'Product, or Service for courier/transport charges.'),
    ('Department / Dept No', 'Department / Dept', 'Unchanged. Some items look mis-filed (e.g. lights under ACRYLIC); not changed.'),
    ('Sub Code', 'SCode', 'Uppercased. Near-duplicates (BA / BAL / BALL) left as they are.'),
    ('Supplier Item No', 'Product No.', 'Stored as text (850043.0 -> 850043).'),
    ('Barcode', 'Supl Code', 'Barcodes Excel had rounded (6910000000000) are blanked and listed in Needs Review.'),
    ('HSN', 'HSN', 'Almost all blank. Needed for GST invoices; fill in before billing goes live.'),
    ('MRP', 'Ret.Rate', 'Unchanged.'),
    ('Offer Price', 'S.Price', 'Blank where the original was 0 (no offer: sell at MRP).'),
    ('Discount %', 'DiscP', 'Recalculated from MRP and Offer Price; wrong originals listed in Needs Review.'),
    ('Cost Price', 'D.Rate', 'Unchanged.'),
    ('GST % (Sale) / Tax % (Purchase)', 'S.Tax% / P.Tax%', 'Unchanged.'),
    ('Manufacturer / Origin', 'Mfr.Name / Mfr.Add.', 'Unchanged.'),
    ('Added Year / Month', 'Mfr.Date', 'Month/year typed as "9/19" had been read by Excel as 19-Sep-2026; converted back to Sep 2019.'),
    ('In Season 2026', '(new, to fill)', 'Put Y for products on sale this season.'),
    ('Source Row', '(new)', 'Row number in the original file, for tracing back.'),
    ('', '', ''),
    ('Source file', SRC.replace('\\', '/').split('/')[-1],
     f'{sh.nrows - 3} data rows read; {len(products)} kept, {len(removed)} removed, {len(review)} flagged for review.'),
]
for n in NOTES:
    notes.append(n)
for c in notes[1]:
    c.font, c.fill = HEAD, FILL
for L, w in zip('ABC', (30, 22, 100)):
    notes.column_dimensions[L].width = w
for rw in notes.iter_rows(min_row=2):
    for c in rw:
        c.alignment = Alignment(wrap_text=True, vertical='top')

wb.save(OUT)
print(f'kept={len(products)} removed={len(removed)} review={len(review)}')
cnt = Counter()
for rv in review:
    for i in rv[4].split('; '):
        cnt[re.sub(r'\(.*', '', i).strip()] += 1
for k, v in cnt.most_common():
    print(f'  {v:5}  {k}')
years = Counter(p['Added Year'] for p in products if p['Added Year'])
print('years:', sorted(years.items()), 'none:', sum(1 for p in products if not p['Added Year']))
