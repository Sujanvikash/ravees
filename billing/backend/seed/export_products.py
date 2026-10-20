"""
Export products-clean.xlsx to the JSON the billing frontend reads, until the billing backend exists.

    python export_products.py products-clean.xlsx ../../frontend/src/data/products.json

Needs: pip install openpyxl. Re-run it after editing products-clean.xlsx (e.g. the "In Season 2026" column).
Cost price, manufacturer and the other purchase details stay out: the counter screen never shows them.
"""
import json
import sys

from openpyxl import load_workbook

SRC, OUT = sys.argv[1], sys.argv[2]
ws = load_workbook(SRC, read_only=True, data_only=True)['Products']
rows = ws.iter_rows(values_only=True)
header = next(rows)
season_col = next(h for h in header if str(h).startswith('In Season'))


def text(v):
    return '' if v is None else str(v).strip()


def money(v):
    return float(v) if isinstance(v, (int, float)) else 0


products = []
for values in rows:
    r = dict(zip(header, values))
    if not text(r['Code']):
        continue
    products.append({
        'code': text(r['Code']),
        'name': text(r['Name']),
        'type': text(r['Type']) or 'Product',
        'dept': text(r['Department']),
        'barcode': text(r['Barcode']),
        'hsn': text(r['HSN']),
        'mrp': money(r['MRP']),
        'offer': money(r['Offer Price']),
        'gst': money(r['GST % (Sale)']),
        'inSeason': text(r[season_col]).upper() in ('Y', 'YES', '1', 'TRUE'),
    })

with open(OUT, 'w', encoding='utf-8') as f:
    json.dump({'season': season_col.replace('In Season', '').strip(), 'products': products},
              f, ensure_ascii=False, separators=(',', ':'))

print(f'{len(products)} products, {sum(p["inSeason"] for p in products)} in season -> {OUT}')
