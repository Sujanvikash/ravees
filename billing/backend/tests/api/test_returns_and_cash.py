import uuid
from io import BytesIO

from openpyxl import load_workbook

from .conftest import login
from .test_sales import sale, season_code, stock_of


def test_return_refunds_cash_and_restocks(client, chennai, admin):
    code = season_code(client, chennai)
    s = client.post('/sales', json=sale([{'code': code, 'qty': 3}]), headers=chennai).json()
    before = stock_of(client, chennai, 'chennai', code)

    found = client.get('/returns/bill', params={'bill_no': s['billNo'].lower()}, headers=chennai).json()
    assert found['open'] and found['lines'][0]['sold'] == 3 and found['lines'][0]['returned'] == 0

    r = client.post(
        '/returns',
        json={'id': str(uuid.uuid4()), 'saleId': s['id'], 'lines': [{'code': code, 'qty': 2}], 'reason': 'Too big'},
        headers=chennai,
    )
    assert r.status_code == 201, r.text
    cn = r.json()['creditNote']
    assert cn['returnNo'].startswith('CN-CHN/') and cn['kind'] == 'return' and cn['billNo'] == s['billNo']
    assert cn['refund'] == cn['amount'] == round(s['lines'][0]['price'] * 2)
    assert r.json()['sale'] is None
    assert stock_of(client, chennai, 'chennai', code) == before + 2

    # only one more can come back
    more = {'id': str(uuid.uuid4()), 'billNo': s['billNo'], 'lines': [{'code': code, 'qty': 2}]}
    assert (
        client.post('/returns', json=more, headers=chennai).json()['detail']
        == f'Only 1 of {code} can still come back on {s["billNo"]}.'
    )
    # another showroom, or the admin, can't take it
    assert (
        client.post(
            '/returns', json={**more, 'lines': [{'code': code, 'qty': 1}]}, headers=login(client, 'goa')
        ).status_code
        == 404
    )
    assert client.post('/returns', json=more, headers=admin).status_code == 403


def test_exchange_takes_the_value_off_the_new_bill(client, chennai):
    code = season_code(client, chennai)
    other = season_code(client, chennai, 1)
    s = client.post('/sales', json=sale([{'code': code, 'qty': 1}]), headers=chennai).json()
    body = {
        'id': str(uuid.uuid4()),
        'saleId': s['id'],
        'lines': [{'code': code, 'qty': 1}],
        'exchange': {'id': str(uuid.uuid4()), 'lines': [{'code': other, 'qty': 3}], 'payment': {'method': 'upi'}},
    }
    r = client.post('/returns', json=body, headers=chennai).json()
    cn, new = r['creditNote'], r['sale']
    assert cn['kind'] == 'exchange' and cn['exchangeBillNo'] == new['billNo']
    credit = min(cn['amount'], new['totals']['payable'])
    assert new['payment']['credit'] == credit and new['payment']['creditNote'] == cn['returnNo']
    assert cn['refund'] == max(cn['amount'] - new['totals']['payable'], 0)
    if new['payment']['method'] == 'upi':
        assert new['payment']['received'] == new['totals']['payable'] - credit
    # sent again: nothing new
    assert client.post('/returns', json=body, headers=chennai).json()['creditNote']['returnNo'] == cn['returnNo']


def test_returns_can_be_switched_off(client, chennai, admin):
    code = season_code(client, chennai)
    s = client.post('/sales', json=sale([{'code': code, 'qty': 1}]), headers=chennai).json()
    assert client.patch('/settings', json={'returnDays': 0}, headers=admin).json()['returnDays'] == 0
    r = client.post(
        '/returns',
        json={'id': str(uuid.uuid4()), 'saleId': s['id'], 'lines': [{'code': code, 'qty': 1}]},
        headers=chennai,
    )
    assert r.status_code == 403
    client.patch('/settings', json={'returnDays': 15}, headers=admin)


def test_cash_close(client, admin):
    mumbai = login(client, 'mumbai')
    code = season_code(client, mumbai)
    day = client.get('/cash/day', headers=mumbai).json()
    client.post('/sales', json=sale([{'code': code, 'qty': 1}], 'cash', received=100000), headers=mumbai)
    after = client.get('/cash/day', headers=mumbai).json()
    assert after['bills'] == day['bills'] + 1 and after['cashSales'] > day['cashSales']

    expected = 500 + after['cashSales'] - after['cashRefunds']
    body = {'opening': 500, 'counted': expected - 50, 'floatNext': 1000, 'denominations': {'500': 2}, 'note': 'Short'}
    assert client.post('/cash/closes', json={**body, 'floatNext': expected + 1}, headers=mumbai).status_code == 422
    c = client.post('/cash/closes', json=body, headers=mumbai)
    assert c.status_code == 201, c.text
    assert c.json()['expected'] == expected and c.json()['difference'] == -50
    assert client.post('/cash/closes', json=body, headers=mumbai).status_code == 409
    assert client.post('/cash/closes', json=body, headers=admin).status_code == 403

    closes = client.get('/cash/closes', headers=admin).json()
    assert any(x['location'] == 'mumbai' and x['difference'] == -50 for x in closes)
    assert client.get('/cash/day?location=mumbai', headers=admin).json()['closed']['difference'] == -50


def test_sales_report_excel(client, admin, chennai):
    r = client.get('/reports/sales.xlsx', headers=admin)
    assert r.status_code == 200 and 'attachment' in r.headers['content-disposition']
    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ['Summary', 'Bills', 'Items', 'Returns', 'Cash']
    assert 'Gross profit' in [c.value for c in wb['Summary'][1]]
    assert wb['Summary'].cell(row=wb['Summary'].max_row, column=1).value == 'All showrooms'

    mine = load_workbook(BytesIO(client.get('/reports/sales.xlsx', headers=chennai).content))
    assert 'Gross profit' not in [c.value for c in mine['Summary'][1]]
    assert {row[2] for row in mine['Bills'].iter_rows(min_row=2, values_only=True)} <= {'Chennai'}
